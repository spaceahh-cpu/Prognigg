"""
Generate a polished cybersecurity engineer reference resume as a Word document.
Usage: python generate_cybersecurity_resume.py
Output: CybersecurityEngineer_Resume_Reference.docx
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy


# ── helpers ──────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = OxmlElement(f"w:{'left' if edge == 'start' else 'right' if edge == 'end' else edge}")
        tag.set(qn("w:val"), kwargs.get(edge, "none"))
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), kwargs.get("color", "auto"))
        tcBorders.append(tag)
    tcPr.append(tcBorders)


def add_horizontal_rule(doc, color="1F4E79"):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    return p


def section_heading(doc, title: str):
    add_horizontal_rule(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    return p


def bullet(doc, text: str, bold_prefix: str = ""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.25)
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        run.font.size = Pt(10)
        p.add_run(text).font.size = Pt(10)
    else:
        p.add_run(text).font.size = Pt(10)
    return p


def job_header(doc, title: str, company_loc: str, dates: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(1)
    r1 = p.add_run(title)
    r1.bold = True
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.add_run(f"  —  {company_loc}").font.size = Pt(10)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(3)
    r2 = p2.add_run(dates)
    r2.italic = True
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(0x60, 0x60, 0x60)


# ── main builder ──────────────────────────────────────────────────────────────

def build_resume():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    # ── Header ──────────────────────────────────────────────────────────────
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_before = Pt(0)
    name_p.paragraph_format.space_after = Pt(2)
    name_run = name_p.add_run("JORDAN ALEX RIVERA")
    name_run.bold = True
    name_run.font.size = Pt(22)
    name_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("Cybersecurity Engineer")
    title_run.font.size = Pt(13)
    title_run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_before = Pt(0)
    contact_p.paragraph_format.space_after = Pt(6)
    contact_run = contact_p.add_run(
        "Austin, TX  |  (512) 555-0192  |  j.rivera@email.com  |  "
        "linkedin.com/in/jordanrivera  |  github.com/jordanrivera-sec"
    )
    contact_run.font.size = Pt(9.5)
    contact_run.font.color.rgb = RGBColor(0x50, 0x50, 0x50)

    # ── Summary ──────────────────────────────────────────────────────────────
    section_heading(doc, "Professional Summary")
    summary = doc.add_paragraph()
    summary.paragraph_format.space_before = Pt(3)
    summary.paragraph_format.space_after = Pt(4)
    summary.add_run(
        "Results-driven Cybersecurity Engineer with 8+ years of experience securing enterprise "
        "environments across finance, healthcare, and SaaS sectors. Deep expertise in threat "
        "detection, penetration testing, cloud security architecture, and incident response. "
        "Reduced mean-time-to-detect (MTTD) by 62% at previous employer by re-architecting "
        "SIEM pipelines. Holds OSCP, CISSP, and AWS Security Specialty certifications. "
        "Passionate about building security programs that scale without slowing engineering velocity."
    ).font.size = Pt(10)

    # ── Technical Skills (table) ─────────────────────────────────────────────
    section_heading(doc, "Technical Skills")
    skills = [
        ("Penetration Testing", "Metasploit, Burp Suite Pro, Cobalt Strike, BloodHound, Impacket, Nmap, Nessus"),
        ("Cloud Security",      "AWS Security Hub, GuardDuty, CloudTrail, Azure Defender, GCP SCC, Terraform, CSPM"),
        ("SIEM / Detection",    "Splunk (SPL), Elastic SIEM, Microsoft Sentinel, Chronicle, YARA, Sigma rules"),
        ("Network Security",    "Wireshark, Zeek, Suricata, pfSense, Cisco ASA, Palo Alto NGFW, Zero Trust (Zscaler)"),
        ("Application Security","OWASP Top 10, SAST (Semgrep, Checkmarx), DAST (OWASP ZAP), SCA (Snyk, Dependabot)"),
        ("Scripting / Dev",     "Python, Bash, PowerShell, Go, SQL, REST APIs, Docker, Kubernetes"),
        ("Frameworks / GRC",    "NIST CSF, MITRE ATT&CK, CIS Controls v8, ISO 27001, SOC 2, HIPAA, PCI-DSS"),
        ("Incident Response",   "Volatility, Autopsy, FTK Imager, TheHive, MISP, Velociraptor"),
    ]
    tbl = doc.add_table(rows=len(skills), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.style = "Table Grid"
    for i, (domain, tools) in enumerate(skills):
        row = tbl.rows[i]
        # Domain cell
        dc = row.cells[0]
        dc.width = Inches(1.9)
        dc.paragraphs[0].clear()
        r = dc.paragraphs[0].add_run(domain)
        r.bold = True
        r.font.size = Pt(9.5)
        dc.paragraphs[0].paragraph_format.space_before = Pt(2)
        dc.paragraphs[0].paragraph_format.space_after = Pt(2)
        set_cell_bg(dc, "EBF3FB")
        # Tools cell
        tc = row.cells[1]
        tc.paragraphs[0].clear()
        tc.paragraphs[0].add_run(tools).font.size = Pt(9.5)
        tc.paragraphs[0].paragraph_format.space_before = Pt(2)
        tc.paragraphs[0].paragraph_format.space_after = Pt(2)

    # ── Experience ───────────────────────────────────────────────────────────
    section_heading(doc, "Professional Experience")

    # Job 1
    job_header(doc, "Senior Cybersecurity Engineer", "FinSecure Corp — Austin, TX", "January 2022 – Present")
    bullets_j1 = [
        "Architected cloud-native SIEM migration from on-prem Splunk to Microsoft Sentinel, cutting licensing costs by $340K/year while improving detection coverage by 45%.",
        "Led red team exercises simulating nation-state TTPs (MITRE ATT&CK T1566, T1078, T1021) across 3 business units; findings drove 18 critical remediations before a regulator audit.",
        "Built custom Python-based threat intel pipeline integrating MISP, VirusTotal, and Shodan, reducing analyst triage time by 38%.",
        "Designed and enforced zero-trust network segmentation model across AWS and Azure multi-cloud, eliminating lateral movement paths found in prior pen tests.",
        "Mentored 4 junior analysts; established SOC runbooks that reduced mean-time-to-respond (MTTR) from 4.2 hours to 47 minutes.",
        "Maintained PCI-DSS Level 1 and SOC 2 Type II compliance posture for card-processing environment.",
    ]
    for b in bullets_j1:
        bullet(doc, b)
    kw = doc.add_paragraph()
    kw.paragraph_format.space_before = Pt(4)
    kw.paragraph_format.left_indent = Inches(0.25)
    kr = kw.add_run("Key Win: ")
    kr.bold = True
    kr.font.size = Pt(10)
    kr.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    kw.add_run(
        "Detected and contained a ransomware intrusion (Conti variant) within 11 minutes of "
        "initial beacon callback — prevented estimated $4.2M in operational losses."
    ).font.size = Pt(10)

    # Job 2
    job_header(doc, "Cybersecurity Engineer II", "MedCloud Health — Denver, CO (Remote)", "March 2019 – December 2021")
    bullets_j2 = [
        "Owned application security program for a HIPAA-regulated SaaS platform serving 2.1M patients; integrated SAST/DAST into CI/CD pipelines (GitHub Actions + Semgrep + OWASP ZAP).",
        "Performed 12+ internal web application pen tests annually; tracked findings through full remediation lifecycle in Jira.",
        "Reduced critical/high CVE backlog from 1,400 to under 80 in 6 months by building SLA-driven vulnerability management workflow with CVSS + EPSS prioritization.",
        "Deployed CrowdStrike Falcon EDR across 3,200 endpoints and tuned detection policies, reducing false-positive alert volume by 71%.",
        "Authored threat models (STRIDE) for 6 new product features, catching 3 design-level authentication flaws before development.",
    ]
    for b in bullets_j2:
        bullet(doc, b)

    # Job 3
    job_header(doc, "Security Analyst", "Accenture Federal Services — Washington, DC", "June 2017 – February 2019")
    bullets_j3 = [
        "Monitored government SOC environment (500K+ events/day) using Splunk; triaged incidents aligned to NIST 800-61 IR lifecycle.",
        "Developed 40+ Splunk correlation rules and dashboards improving detection of privilege escalation and data exfiltration patterns.",
        "Supported FedRAMP readiness assessment; produced System Security Plan (SSP) and contributed to continuous monitoring strategy.",
        "Conducted phishing simulations using GoPhish; executive-level reports drove 3× budget increase for security awareness training.",
    ]
    for b in bullets_j3:
        bullet(doc, b)

    # ── Certifications ───────────────────────────────────────────────────────
    section_heading(doc, "Certifications")
    certs = [
        ("Offensive Security Certified Professional (OSCP)", "Offensive Security", "2021"),
        ("Certified Information Systems Security Professional (CISSP)", "(ISC)²", "2022"),
        ("AWS Certified Security – Specialty", "Amazon Web Services", "2023"),
        ("CompTIA Security+ (CE)", "CompTIA", "2017"),
        ("GIAC Certified Incident Handler (GCIH)", "SANS / GIAC", "2020"),
    ]
    ctbl = doc.add_table(rows=1 + len(certs), cols=3)
    ctbl.style = "Table Grid"
    headers = ["Certification", "Issuing Body", "Year"]
    for i, h in enumerate(headers):
        c = ctbl.rows[0].cells[i]
        c.paragraphs[0].clear()
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(9.5)
        set_cell_bg(c, "1F4E79")
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        c.paragraphs[0].paragraph_format.space_before = Pt(2)
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
    for i, (cert, body, year) in enumerate(certs):
        row = ctbl.rows[i + 1]
        for j, val in enumerate((cert, body, year)):
            c = row.cells[j]
            c.paragraphs[0].clear()
            c.paragraphs[0].add_run(val).font.size = Pt(9.5)
            c.paragraphs[0].paragraph_format.space_before = Pt(2)
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
            if i % 2 == 0:
                set_cell_bg(c, "F5F8FF")

    # ── Education ────────────────────────────────────────────────────────────
    section_heading(doc, "Education")
    ep = doc.add_paragraph()
    ep.paragraph_format.space_before = Pt(3)
    ep.paragraph_format.space_after = Pt(1)
    er = ep.add_run("B.S. Computer Science, Concentration: Information Assurance")
    er.bold = True
    er.font.size = Pt(10)
    doc.add_paragraph(
        "University of Texas at San Antonio  |  2017  |  GPA: 3.7  |  "
        "NSA/DHS Center of Academic Excellence in Cyber Defense"
    ).runs[0].font.size = Pt(10)

    # ── Projects & Research ──────────────────────────────────────────────────
    section_heading(doc, "Notable Projects & Research")
    projects = [
        ("CloudHawk (open source)",
         "Python tool for continuous auditing of AWS IAM misconfigurations; 800+ GitHub stars. "
         "Integrates with AWS Security Hub and sends Slack/PagerDuty alerts."),
        ("MITRE ATT&CK Detection Lab",
         "Home lab with Elastic SIEM + Caldera adversary emulation; 35 detection rules documented "
         "and shared publicly (2,500+ monthly blog readers)."),
        ("CVE-2023-XXXX",
         "Responsibly disclosed an SSRF vulnerability in a widely-used open-source API gateway; "
         "awarded $5,000 bug bounty."),
        ("BSides Austin 2023 — Speaker",
         "\"Hunting Lateral Movement in Cloud Logs Without Going Broke on Storage\""),
    ]
    for name, desc in projects:
        bullet(doc, desc, bold_prefix=f"{name}: ")

    # ── Save ─────────────────────────────────────────────────────────────────
    output = "CybersecurityEngineer_Resume_Reference.docx"
    doc.save(output)
    print(f"Resume saved → {output}")


if __name__ == "__main__":
    build_resume()
